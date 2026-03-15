"""
多模态 AI 互动式教学智能体 — 最终版 UI（工业级视觉）
"""
from __future__ import annotations
import json, traceback
from pathlib import Path
from typing import List, Optional
import gradio as gr
from .service import RAGService

# ═══════════════════════════════════════════
# 工业级自定义CSS（Gemini风格 + 功能保留）
# ═══════════════════════════════════════════
CUSTOM_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Nunito:wght@400;600;800&display=swap');
.gradio-container {
    font-family: 'Nunito', 'PingFang SC', sans-serif!important;
    background: linear-gradient(135deg, #eef2f3 0%, #cbd5e1 100%)!important;
}
.glass-panel {
    background: rgba(255, 255, 255, 0.6)!important;
    backdrop-filter: blur(16px)!important;
    border: 1px solid rgba(255, 255, 255, 0.8)!important;
    border-radius: 24px!important;
    box-shadow: 0 10px 40px -10px rgba(30, 41, 59, 0.1)!important;
    padding: 20px!important;
}
button.primary {
    background-color: #58cc02!important; color: white!important;
    font-weight: 800!important; font-size: 1.1rem!important;
    border-radius: 16px!important; border: none!important;
    box-shadow: 0 4px 0 #58a700!important;
    transition: all 0.15s cubic-bezier(0.4, 0, 0.2, 1)!important;
    text-transform: uppercase; letter-spacing: 1px;
}
button.primary:hover { background-color: #61df02!important; transform: translateY(-1px)!important; box-shadow: 0 5px 0 #58a700!important; }
button.primary:active { transform: translateY(4px)!important; box-shadow: 0 0 0 #58a700!important; }
button.secondary {
    background-color: #ffffff!important; color: #64748b!important;
    font-weight: 700!important; border-radius: 16px!important;
    border: 2px solid #e2e8f0!important; box-shadow: 0 4px 0 #e2e8f0!important;
}
button.secondary:active { transform: translateY(4px)!important; box-shadow: none!important; }
.preview-html {
    border-radius: 24px!important; overflow: hidden!important;
    background: white!important; box-shadow: 0 15px 35px rgba(0,0,0,0.08)!important;
    border: 2px solid #f1f5f9!important; min-height: 500px;
}
.tabs > div > button { font-weight: 700!important; font-size: 1.05rem!important; }
"""

_empty = """
<div style='display:flex;flex-direction:column;align-items:center;justify-content:center;height:100%;min-height:500px;color:#94a3b8;font-family:Nunito,sans-serif;'>
    <svg width="80" height="80" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.5"><rect x="2" y="3" width="20" height="14" rx="2"></rect><line x1="8" y1="21" x2="16" y2="21"></line><line x1="12" y1="17" x2="12" y2="21"></line></svg>
    <h3 style="margin:10px 0 0;font-size:1.5rem;font-weight:800;color:#475569;">引擎就绪</h3>
    <p style="margin-top:12px;font-size:1rem;max-width:300px;text-align:center;line-height:1.6;">在左侧输入教学需求，AI将自动生成课件。</p>
</div>
"""

def build_ui(cfg_path: str = "config.yaml") -> gr.Blocks:
    service = RAGService(cfg_path)
    dialogue_mgr = None
    try:
        from .dialogue.dialogue_manager import DialogueManager
        dialogue_mgr = DialogueManager(cfg_path)
    except Exception: pass
    pipeline = None
    try:
        from .courseware_pipeline import CoursewarePipeline
        pipeline = CoursewarePipeline(cfg_path)
    except Exception: pass
    try: idxs = service.list_indexes() or ["kb"]
    except: idxs = ["kb"]

    # ── 工具函数 ──
    def asr_to_text(audio_path) -> str:
        if not audio_path: return ""
        try:
            from .tools.asr import transcribe_audio
            r = transcribe_audio(str(audio_path))
            return str(r) if not isinstance(r, str) else r
        except Exception as e:
            return f"[语音识别需安装faster-whisper: {e}]"

    def pptx_preview(path):
        if not path or not Path(path).exists(): return _empty
        try:
            from .ppt_preview import pptx_preview_html
            return pptx_preview_html(path, max_pages=12)
        except:
            try:
                from pptx import Presentation
                prs = Presentation(path)
                lines = [f"<p><b>第{i+1}页:</b> {' | '.join(s.text.strip()[:50] for s in sl.shapes if hasattr(s,'text') and s.text.strip())[:3]}</p>" for i,sl in enumerate(prs.slides)]
                return "<div>" + "\n".join(lines) + "</div>"
            except Exception as e: return f"<p>预览失败: {e}</p>"

    def docx_preview(path):
        if not path or not Path(path).exists(): return "无Word文件"
        try:
            import docx; doc = docx.Document(path)
            lines = []
            for p in doc.paragraphs[:30]:
                t = p.text.strip()
                if not t: continue
                sn = p.style.name if p.style else ""
                if "Heading" in sn: lines.append(f"\n## {t}")
                else: lines.append(t)
            return "\n".join(lines)[:3000]
        except Exception as e: return f"预览失败: {e}"

    # ── 对话 ──
    def chat(user_msg, audio, history, files, note):
        if audio and not user_msg.strip(): user_msg = asr_to_text(audio)
        if not user_msg.strip(): return history or [], "", "请输入文字或语音", None
        if not dialogue_mgr: return (history or [])+[[user_msg,"⚠️ 对话模块未加载"]], "", "未就绪", None
        try:
            fps = [f.name if hasattr(f,'name') else str(f) for f in (files if isinstance(files,list) else [files])] if files else []
            ns = [note]*len(fps) if fps else None
            reply, state = dialogue_mgr.chat(user_input=user_msg, uploaded_files=fps or None, file_notes=ns)
            ch = (history or [])+[[user_msg, reply]]
            info = f"阶段: {state.value}\n已收集:\n{dialogue_mgr.collected.to_text()}\n\n缺失: {', '.join(dialogue_mgr.collected.missing_fields()) or '✅ 信息完整'}"
            return ch, "", info, None
        except Exception as e:
            return (history or [])+[[user_msg,f"❌ {e}"]], "", str(e), None

    def reset():
        if dialogue_mgr: dialogue_mgr.reset()
        return [], "", "已重置", None

    # ── 一键生成 ──
    def gen_all(_st):
        if not pipeline:
            return "❌ 流水线未加载", _empty, "未生成", _empty, _empty, None, None, None, None, None
        if dialogue_mgr and not dialogue_mgr.collected.topic and not dialogue_mgr.history:
            return "⚠️ 请先在左侧对话中描述教学需求，然后再点此按钮。", _empty, "", _empty, _empty, None, None, None, None, None
        try:
            r = pipeline.generate_all(dialogue_mgr=dialogue_mgr, rag_service=service, output_types=["ppt","docx","game"])
            st = r.summary()
            ph = pptx_preview(r.pptx_path) if r.pptx_path else _empty
            dt = docx_preview(r.docx_path) if r.docx_path else "未生成"
            gh = r.game_html or _empty
            ah = r.animation_html or _empty
            pf = r.pptx_path if r.pptx_path and Path(r.pptx_path).exists() else None
            df = r.docx_path if r.docx_path and Path(r.docx_path).exists() else None
            gf = r.game_path if r.game_path and Path(r.game_path).exists() else None
            af = r.animation_path if r.animation_path and Path(r.animation_path).exists() else None
            return st, ph, dt, gh, ah, pf, df, gf, af, r
        except Exception as e:
            return f"❌ {e}\n{traceback.format_exc()}", _empty, "", _empty, _empty, None, None, None, None, None

    # ── 迭代修改 ──
    def regen(fb, rp, rd, rg, prev):
        if not pipeline: return "❌ 未加载", _empty, "", _empty, _empty, None, None, None, None, prev
        if not fb.strip(): return "请输入修改意见", _empty, "", _empty, _empty, None, None, None, None, prev
        if not prev: return "请先生成课件", _empty, "", _empty, _empty, None, None, None, None, None
        rt = []
        if rp: rt.append("ppt")
        if rd: rt.append("docx")
        if rg: rt.append("game")
        if not rt: rt = ["ppt","docx","game"]
        try:
            r = pipeline.regenerate_with_feedback(feedback=fb, previous_result=prev, rag_service=service, regenerate_types=rt)
            st = "🔄 修改完成:\n"+r.summary()
            return st, pptx_preview(r.pptx_path) if r.pptx_path else "", docx_preview(r.docx_path) if r.docx_path else "", r.game_html or "", r.animation_html or getattr(prev,'animation_html','') or "", r.pptx_path, r.docx_path, r.game_path, getattr(r,'animation_path',None) or getattr(prev,'animation_path',None), r
        except Exception as e:
            return f"❌ {e}", "", "", "", "", None, None, None, None, prev

    # ── RAG ──
    def query(q, idx, k, tr):
        try:
            res = service.query(question=q, indexes=idx or ["kb"], top_k=int(k), evidence_tag="ui", enable_trace=True)
            hv = res.get("human_view",{}) or {}
            return hv.get("answer_md",""), hv.get("evidence_md",""), hv.get("sources_md",""), hv.get("trace_md","") if tr else "", json.dumps({"answer":res.get("answer")},ensure_ascii=False,indent=2)
        except Exception as e: return "","","",f"错误:{e}",""

    # ═══════════════════════════════════════════
    # BUILD UI（工业级视觉 + 完整后端连接）
    # ═══════════════════════════════════════════
    theme = gr.themes.Soft(
        primary_hue="indigo", secondary_hue="blue", neutral_hue="slate",
        font=[gr.themes.GoogleFont("Nunito"), "Arial", "sans-serif"]
    )

    with gr.Blocks(theme=theme, css=CUSTOM_CSS, title="多模态AI互动式教学智能体") as demo:
        cw = gr.State(None)

        gr.Markdown(
            """<div style="display:flex;align-items:center;gap:15px;padding:15px;">
            <div style="background:linear-gradient(135deg,#6366f1,#8b5cf6);padding:12px;border-radius:16px;color:white;">
                <svg width="32" height="32" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.5"><path d="M12 2L2 7l10 5 10-5-10-5z"></path><path d="M2 17l10 5 10-5"></path><path d="M2 12l10 5 10-5"></path></svg>
            </div>
            <div>
                <h1 style="margin:0;font-size:2rem;font-weight:800;color:#1e293b;">多模态 AI 互动式教学智能体</h1>
                <p style="color:#64748b;font-size:1rem;margin:5px 0 0;font-weight:600;">A04 锐捷网络 · 对话→生成PPT+Word+游戏+动画→预览→修改→再生成</p>
            </div></div>"""
        )

        # ══════ Tab 1: 教学对话 & 课件生成 ══════
        with gr.Tab("✨ 智能共创中心"):
            with gr.Row():
                with gr.Column(scale=4, elem_classes="glass-panel"):
                    gr.Markdown("### 🎙️ 描述教学需求")
                    chatbot = gr.Chatbot(height=300, show_label=False)
                    with gr.Row():
                        mi = gr.Textbox(show_label=False, placeholder="描述需求，如：帮我生成关于DNA转录的互动课...", scale=5)
                        sb = gr.Button("发送 ✈️", elem_classes="secondary", scale=1)
                    au = gr.Audio(sources=["microphone"], type="filepath", label="🎙️ 语音输入")
                    with gr.Accordion("📎 上传参考资料", open=False):
                        gr.Markdown("<span style='color:#64748b;font-size:0.9em;'>上传PDF/Word/PPT/音视频，AI从中提取知识。每次上传请说明用途。</span>")
                        rf = gr.File(label="拖拽文件至此", file_count="multiple")
                        rn = gr.Textbox(label="资料用途说明", placeholder="如：参照此PDF第3章的格式和知识点")
                    si = gr.Textbox(label="💡 已收集信息", interactive=False, lines=3)
                    with gr.Row():
                        rb = gr.Button("🔄 重置对话", elem_classes="secondary")
                        gb = gr.Button("🚀 一键生成全部课件", variant="primary", elem_classes="primary")
                    gs = gr.Textbox(label="⏳ 生成状态", interactive=False, lines=3)

                with gr.Column(scale=7):
                    with gr.Tabs():
                        with gr.Tab("🎮 互动游戏"):
                            gp = gr.HTML(value=_empty, elem_classes="preview-html")
                            gd = gr.File(label="⏬ 下载游戏 (.html)")
                        with gr.Tab("🎬 知识动画"):
                            ap = gr.HTML(value=_empty, elem_classes="preview-html")
                            ad = gr.File(label="⏬ 下载动画 (.html)")
                        with gr.Tab("📊 PPT课件"):
                            pp = gr.HTML(value=_empty, elem_classes="preview-html")
                            pd = gr.File(label="⏬ 下载PPT (.pptx)")
                        with gr.Tab("📝 Word教案"):
                            dp = gr.Textbox(label="教案预览", lines=18, interactive=False, value="等待生成")
                            dd = gr.File(label="⏬ 下载教案 (.docx)")

            sb.click(fn=chat, inputs=[mi,au,chatbot,rf,rn], outputs=[chatbot,mi,si,au])
            mi.submit(fn=chat, inputs=[mi,au,chatbot,rf,rn], outputs=[chatbot,mi,si,au])
            rb.click(fn=reset, outputs=[chatbot,mi,si,au])
            gb.click(fn=gen_all, inputs=[cw], outputs=[gs,pp,dp,gp,ap,pd,dd,gd,ad,cw])

        # ══════ Tab 2: 迭代修改 ══════
        with gr.Tab("✏️ 迭代微调"):
            with gr.Row():
                with gr.Column(scale=4, elem_classes="glass-panel"):
                    gr.Markdown("### 🎯 输入修改意见")
                    fi = gr.Textbox(label="修改要求", placeholder="如：把第3页简化 / 游戏改排序题 / 教案增加讨论环节", lines=5)
                    with gr.Row():
                        fp = gr.Checkbox(value=True, label="PPT")
                        fd = gr.Checkbox(value=True, label="Word")
                        fg = gr.Checkbox(value=True, label="游戏")
                    fb_btn = gr.Button("✨ 应用修改并重新生成", variant="primary", elem_classes="primary")
                    fs = gr.Textbox(label="状态", lines=3, interactive=False)
                with gr.Column(scale=7):
                    with gr.Tabs():
                        with gr.Tab("🎮 游戏"):
                            rgp = gr.HTML(value=_empty, elem_classes="preview-html")
                            rgd = gr.File(label="下载游戏")
                        with gr.Tab("🎬 动画"):
                            rap = gr.HTML(value=_empty, elem_classes="preview-html")
                            rad = gr.File(label="下载动画")
                        with gr.Tab("📊 PPT"):
                            rpp = gr.HTML(value=_empty, elem_classes="preview-html")
                            rpd = gr.File(label="下载PPT")
                        with gr.Tab("📝 教案"):
                            rdp = gr.Textbox(lines=15, interactive=False)
                            rdd = gr.File(label="下载教案")

            fb_btn.click(fn=regen, inputs=[fi,fp,fd,fg,cw], outputs=[fs,rpp,rdp,rgp,rap,rpd,rdd,rgd,rad,cw])

        # ══════ Tab 3: 知识问答 ══════
        with gr.Tab("🔍 知识库问答"):
            with gr.Column(elem_classes="glass-panel"):
                q = gr.Textbox(label="提问", placeholder="向知识库提问...", lines=2)
                with gr.Row():
                    qi = gr.Dropdown(choices=idxs, value=["kb"] if "kb" in idxs else idxs[:1], multiselect=True, label="索引")
                    qk = gr.Slider(1,20,value=5,step=1,label="Top-K")
                    qt = gr.Checkbox(value=False, label="显示Trace")
                qb = gr.Button("查询", variant="primary")
                qa = gr.Markdown()
                with gr.Accordion("📄 证据来源", open=False):
                    qe = gr.Markdown(); qs_md = gr.Markdown()
                with gr.Accordion("⚙️ Trace日志", open=False):
                    qtr = gr.Markdown()
                with gr.Accordion("JSON", open=False):
                    qj = gr.Code(language="json")
            qb.click(fn=query, inputs=[q,qi,qk,qt], outputs=[qa,qe,qs_md,qtr,qj])

        # ══════ Tab 4: 知识库管理 ══════
        with gr.Tab("📚 知识库管理"):
            with gr.Column(elem_classes="glass-panel"):
                ks = gr.Textbox(label="索引列表", interactive=False)
                gr.Button("🔄 刷新").click(fn=lambda: "\n".join(f"  - {x}" for x in service.list_indexes()), outputs=[ks])
                gr.Markdown("---")
                with gr.Row():
                    id_ = gr.Textbox(label="资料目录", value="knowledge_base")
                    ii = gr.Textbox(label="索引名", value="kb")
                    iis = gr.Textbox(label="Session", value="kb")
                    ir = gr.Checkbox(value=True, label="重建索引")
                ib = gr.Button("📥 开始入库", variant="primary", elem_classes="primary")
                io = gr.Textbox(label="结果", lines=3, interactive=False)
                def do_ingest(d,i,s,r):
                    try:
                        if r: service.engine.reset_index(i)
                        return f"✅ {service.engine.ingest_dir(d,index=i,source_type='knowledge_base',session_id=s or 'kb')}"
                    except Exception as e: return f"❌ {e}"
                ib.click(fn=do_ingest, inputs=[id_,ii,iis,ir], outputs=[io])

    return demo

def main():
    demo = build_ui("config.yaml"); demo.queue(); demo.launch(server_name="0.0.0.0", server_port=7860, share=False)

if __name__ == "__main__": main()

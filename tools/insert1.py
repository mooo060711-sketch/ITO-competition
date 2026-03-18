import os
import io
import logging
from PIL import Image
from pptx import Presentation

# 配置日志
logger = logging.getLogger("ppt_injection")

def get_visual_sorted_pictures(slide):
    """
    将幻灯片中的图片按视觉习惯排序。
    优化：动态计算行容差，防止不同尺寸 PPT 导致排序错乱。
    """
    pictures = [s for s in slide.shapes if s.shape_type in (13, 14)]
    
    # 🌟 动态容差：取幻灯片高度的 1% 作为行判定基准
    slide_height = slide.part.package.presentation_part.presentation.slide_height
    ROW_THRESHOLD = slide_height // 100 
    
    pictures.sort(key=lambda p: (p.top // ROW_THRESHOLD, p.left))
    return pictures

def replace_image_by_coordinate_logic(ppt_filepath: str, coordinate_mapping: dict, output_path: str = None):
    """
    V11.0 强化版：支持自定义输出路径，增强格式兼容性。
    """
    if not coordinate_mapping:
        return ppt_filepath

    try:
        prs = Presentation(ppt_filepath)
    except Exception as e:
        logger.error(f"无法打开 PPT 文件 {ppt_filepath}: {e}")
        return ""
    
    for page_str, assignments in coordinate_mapping.items():
        try:
            page_num = int(page_str)
            slide = prs.slides[page_num - 1]
            sorted_pics = get_visual_sorted_pictures(slide)
            
            for rank, new_img_path in assignments.items():
                rank_idx = int(rank) - 1
                if rank_idx < 0 or rank_idx >= len(sorted_pics):
                    continue
                
                if not os.path.exists(new_img_path):
                    logger.warning(f"图片不存在: {new_img_path}")
                    continue

                target_pic = sorted_pics[rank_idx]
                
                # 执行二进制注入
                try:
                    blip = target_pic._element.blipFill.blip
                    NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
                    rId = blip.get(f"{NS_R}embed") or blip.get(f"{NS_R}link")
                    
                    image_part = slide.part.related_part(rId)
                    orig_content_type = image_part.content_type 
                    
                    # 图像格式安全处理
                    img = Image.open(new_img_path)
                    if 'jpeg' in orig_content_type and img.mode in ('RGBA', 'P', 'LA'):
                        img = img.convert('RGB') # 强制转RGB解决透明背景存为JPG报错
                    
                    img_byte_arr = io.BytesIO()
                    save_format = 'JPEG' if 'jpeg' in orig_content_type else 'PNG'
                    img.save(img_byte_arr, format=save_format)
                    
                    # 🌟 注入 Blob
                    image_part._blob = img_byte_arr.getvalue()
                    logger.info(f"✅ 第{page_num}页位置[{rank}] 替换成功")
                except Exception as e:
                    logger.error(f"❌ 注入失败: {e}")
        except Exception:
            continue

    # 保存逻辑优化
    if not output_path:
        output_path = ppt_filepath.replace(".pptx", "_final.pptx")
    
    prs.save(output_path)
    return output_path

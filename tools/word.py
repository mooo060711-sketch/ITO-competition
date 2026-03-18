import os
import io
import re  # 🌟 核心新增：用于正则清洗加粗标记
import logging
import time
from pptx import Presentation
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser

# ===================== 0. 配置与初始化 =====================
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger("lesson_plan_engine")
load_dotenv()

from src.llm_config import load_llm_config

config = load_llm_config()

def init_lesson_llm():
    return ChatOpenAI(
        model=config['models']['lesson'],  # deepseek-chat
        api_key=config['api_key'],
        base_url=config['base_url'],
        temperature=0.3,
        timeout=120
    )
parser = StrOutputParser()

# ===================== 1. 反向内容提取 =====================
def extract_text_from_ppt(ppt_filepath: str) -> str:
    """逆向提取 PPT 文字内容"""
    logger.info(f"🔍 正在执行内容反写，解析 PPT: {ppt_filepath}")
    
    try:
        prs = Presentation(ppt_filepath)
    except Exception as e:
        logger.error(f"❌ PPT 解析失败: {e}")
        return ""

    all_slides_text = []
    for i, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        
        page_content = f"【第 {i+1} 页幻灯片内容】\n" + "\n".join(slide_texts)
        all_slides_text.append(page_content)
        
    return "\n\n".join(all_slides_text)

# ===================== 2. 教案生成 Prompt 定义 =====================
# 🌟 优化：增加了严禁输出引用标记的强制指令
LESSON_PLAN_PROMPT = """
你是一名深耕一线、具备深厚教育心理学背景的生物金牌教研专家。
请根据提供的【教学意图】与【PPT真实内容】，编写一份 100% 契合且具备教学深度的配套教案。

# 一、 课题名称
# 二、 教学目标与核心素养
# 三、 教学重难点
# 四、 课时安排
# 五、 教学过程
   - 请逐页对照 PPT 内容，按以下模版展开：
   - a) 教师活动：
     * 【知识深度拆解】：PPT 上的要点必须全部出现，以 `->` 开头进行深度拓展和启发式提问。
   - b) 学生活动：学生观察、思考、讨论的具体内容。
   - c) 设计意图：体现核心素养。
# 六、 教学方法与板书设计
# 七、 课后作业（基础、提升、挑战三层，每层 3 道题）

【输出禁令】：
1. 严禁包含 ```markdown 标记。
2. 🌟 严禁在正文中输出任何形如 或 [1] 的引用/溯源标记。
3. 严禁输出任何关于大纲生成的元数据。

---
【全量教学意图要素】：
{intent_info}

【受众属性】：
{audience_info}

【PPT 真实排版内容】：
{real_ppt_content}
"""

lesson_prompt = ChatPromptTemplate.from_messages([
    ("system", "你是一个资深的教案编写专家。"),
    ("human", LESSON_PLAN_PROMPT)
])

# ===================== 3. 格式转换：Markdown 转 Word =====================
def markdown_to_docx(md_text: str, save_path: str):
    """将生成的 Markdown 教案转换为带排版格式的 Word 文档"""
    doc = Document()
    
    # 设置全文字体
    doc.styles['Normal'].font.name = u'宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    doc.styles['Normal'].font.size = Pt(12)

    # 预处理：移除干扰字符
    md_text = md_text.replace('```markdown', '').replace('```', '').strip()
    
    lines = md_text.split('\n')
    for line in lines:
        raw_line = line.strip()
        
        # 🌟 修复 1：保留空行，增加 Word 文档的呼吸感
        if not raw_line:
            doc.add_paragraph("") 
            continue
            
        # 🌟 修复 2：增强多级标题处理（支持到三级）
        if raw_line.startswith('# '):
            doc.add_heading(raw_line[2:].replace('**', ''), level=1)
        elif raw_line.startswith('## '):
            doc.add_heading(raw_line[3:].replace('**', ''), level=2)
        elif raw_line.startswith('### '):
            doc.add_heading(raw_line[4:].replace('**', ''), level=3)
        
        # 🌟 修复 3：处理列表并清洗加粗符号
        elif raw_line.startswith('- ') or raw_line.startswith('* '):
            clean_text = raw_line[2:].replace('**', '')
            doc.add_paragraph(clean_text, style='List Bullet')
        
        # 🌟 修复 4：处理正文并移除 Markdown 加粗标记
        else:
            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', raw_line)
            doc.add_paragraph(clean_text)

    doc.save(save_path)
    logger.info(f"✅ 教案物理文件已优化并保存: {save_path}")

# ===================== 4. 主调度函数 =====================
def generate_word_lesson_plan_from_ppt(intent_slots, ppt_filepath: str, output_dir: str = "./") -> dict:
    """教案生成主调度"""
    # 1. 内容反写
    real_ppt_content = extract_text_from_ppt(ppt_filepath)
    if not real_ppt_content:
        logger.error("❌ 无法从 PPT 中提取内容，教案中止")
        return None
        
    # 2. 组装输入要素
    intent_info = (
        f"教学主题：{intent_slots.theme}\n"
        f"核心目标：{intent_slots.teaching_objectives}\n"
        f"重点难点：{intent_slots.key_points}\n"
        f"必含知识点：{intent_slots.required_knowledge}\n"
        f"特殊要求：{intent_slots.other_requirements}"
    )
    
    # 3. LLM 撰写
    logger.info("📝 正在撰写深度教案内容...")
    try:
        llm_chain = lesson_prompt | init_lesson_llm() | parser
        lesson_plan_md = llm_chain.invoke({
            "intent_info": intent_info,
            "audience_info": intent_slots.audience,
            "real_ppt_content": real_ppt_content
        })
        
        # 4. 生成 Word
        if not os.path.exists(output_dir): 
            os.makedirs(output_dir)
            
        safe_theme = intent_slots.theme.replace(" ", "_").replace("/", "_")
        docx_filename = f"配套教案_{safe_theme}_{int(time.time())}.docx"
        docx_filepath = os.path.join(output_dir, docx_filename)
        
        markdown_to_docx(lesson_plan_md, docx_filepath)
        
        return {
            "lesson_plan_str": lesson_plan_md,
            "docx_filepath": docx_filepath
        }
        
    except Exception as e:
        logger.error(f"❌ 教案生成流程中断: {e}")
        return None
